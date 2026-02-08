"""
Migration Script: Move to new watch assets and data

1. Extract zip files from newer-stuff/
2. Convert Excel tabs to CSVs (same format as existing)
3. Move old watch assets to legacy/
4. Keep Sunglasses unchanged
"""

import os
import shutil
import zipfile
import openpyxl
import csv
from pathlib import Path

# Paths
ROOT = Path(__file__).parent.parent
NEWER_STUFF = ROOT / "newer-stuff"
PRODUCTS_ASSETS = ROOT / "products-assets"
PRODUCTS_CSV = ROOT / "products-csv"
LEGACY = ROOT / "legacy"

EXCEL_FILE = NEWER_STUFF / "Watches - Men, Women, Kids.xlsx"

# Zip files and their target extraction folders
ZIP_MAPPINGS = {
    "Men Watches Images.zip": "Men Watches Images",
    "Women Watch Images.zip": "Women Watch Images",
    "Kids Watches.zip": "Kids Watches",
}

# Sheet mappings: sheet_name -> (csv_filename, gender)
SHEET_MAPPINGS = {
    "Men Watch": ("products_Men Watches.csv", "MEN"),
    "Women Watch": ("products_Women Watch.csv", "WOMEN"),
    "Kids Watch": ("products_Kids Watch.csv", "KIDS"),
}


def move_old_assets_to_legacy():
    """Move current watch asset folders to legacy"""
    print("\n=== Moving old watch assets to legacy ===")
    
    folders_to_move = ["Men Watches Images", "Women Watch Images", "Kids Watches"]
    legacy_assets = LEGACY / "old-watch-assets"
    legacy_assets.mkdir(parents=True, exist_ok=True)
    
    for folder in folders_to_move:
        src = PRODUCTS_ASSETS / folder
        dst = legacy_assets / folder
        if src.exists():
            if dst.exists():
                shutil.rmtree(dst)
            print(f"  Moving {folder} -> legacy/old-watch-assets/")
            shutil.move(str(src), str(dst))
        else:
            print(f"  Skipping {folder} (not found)")


def move_old_csvs_to_legacy():
    """Move current watch CSVs to legacy"""
    print("\n=== Moving old watch CSVs to legacy ===")
    
    legacy_csv = LEGACY / "old-watch-csvs"
    legacy_csv.mkdir(parents=True, exist_ok=True)
    
    for sheet_name, (csv_filename, _) in SHEET_MAPPINGS.items():
        src = PRODUCTS_CSV / csv_filename
        dst = legacy_csv / csv_filename
        if src.exists():
            print(f"  Moving {csv_filename} -> legacy/old-watch-csvs/")
            shutil.move(str(src), str(dst))


def extract_zips():
    """Extract zip files to products-assets"""
    print("\n=== Extracting new asset zips ===")
    
    for zip_name, target_folder in ZIP_MAPPINGS.items():
        zip_path = NEWER_STUFF / zip_name
        extract_to = PRODUCTS_ASSETS / target_folder
        
        if not zip_path.exists():
            print(f"  Warning: {zip_name} not found!")
            continue
            
        print(f"  Extracting {zip_name}...")
        
        # Remove existing folder if exists
        if extract_to.exists():
            shutil.rmtree(extract_to)
        
        # Extract zip
        with zipfile.ZipFile(zip_path, 'r') as zf:
            # Get the top-level folder name in the zip
            names = zf.namelist()
            top_folders = set()
            for name in names:
                parts = name.split('/')
                if parts[0]:
                    top_folders.add(parts[0])
            
            # Extract to temp location first
            temp_extract = PRODUCTS_ASSETS / "_temp_extract"
            if temp_extract.exists():
                shutil.rmtree(temp_extract)
            
            zf.extractall(temp_extract)
            
            # If there's a single top-level folder, move its contents
            extracted_items = list(temp_extract.iterdir())
            if len(extracted_items) == 1 and extracted_items[0].is_dir():
                # Single folder inside zip - rename it to target
                shutil.move(str(extracted_items[0]), str(extract_to))
                shutil.rmtree(temp_extract)
            else:
                # Multiple items - rename temp folder to target
                shutil.move(str(temp_extract), str(extract_to))
        
        print(f"    -> Extracted to {target_folder}/")


def convert_excel_to_csvs():
    """Convert Excel tabs to CSV files"""
    print("\n=== Converting Excel to CSVs ===")
    
    if not EXCEL_FILE.exists():
        print(f"  Error: Excel file not found: {EXCEL_FILE}")
        return
    
    # Load with data_only=True to get computed values
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    for sheet_name, (csv_filename, gender) in SHEET_MAPPINGS.items():
        if sheet_name not in wb.sheetnames:
            print(f"  Warning: Sheet '{sheet_name}' not found!")
            continue
        
        sheet = wb[sheet_name]
        csv_path = PRODUCTS_CSV / csv_filename
        
        print(f"  Processing {sheet_name} -> {csv_filename}")
        
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:  # Skip header
                continue
            
            # Check if row has data (NO. column is not None)
            if row[0] is None:
                continue
            
            # Extract: Image SKU (for folder), SKU, Price, and combine details for description
            image_sku = row[1]  # Image Name SKU column - maps to asset folder
            sku = row[2]  # SKU column
            price = row[4]  # Price column
            
            # Combine detail columns (6-14, indices 6-14) filtering out ' I ' separators
            details = []
            for j in range(6, 15):
                if j < len(row) and row[j] is not None:
                    val = str(row[j]).strip()
                    if val and val != 'I' and val != ' I ':
                        details.append(val)
            
            desc = " ; ".join(details)
            
            rows.append({
                'id': sku,
                'imageSku': image_sku,
                'category': 'WATCHES',
                'gender': gender,
                'price': price,
                'desc': desc
            })
        
        # Write CSV
        with open(csv_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=['id', 'imageSku', 'category', 'gender', 'price', 'desc'])
            writer.writeheader()
            writer.writerows(rows)
        
        print(f"    -> Wrote {len(rows)} products")


def list_extracted_structure():
    """Show what we extracted"""
    print("\n=== New asset structure ===")
    for folder in ["Men Watches Images", "Women Watch Images", "Kids Watches"]:
        folder_path = PRODUCTS_ASSETS / folder
        if folder_path.exists():
            items = list(folder_path.iterdir())
            print(f"  {folder}/: {len(items)} items")
            # Show first few
            for item in items[:3]:
                print(f"    - {item.name}/")
            if len(items) > 3:
                print(f"    ... and {len(items) - 3} more")


def main():
    print("=" * 60)
    print("MIGRATION: New Watch Assets and Data")
    print("=" * 60)
    
    # Step 1: Move old assets to legacy
    move_old_assets_to_legacy()
    
    # Step 2: Move old CSVs to legacy
    move_old_csvs_to_legacy()
    
    # Step 3: Extract new zips
    extract_zips()
    
    # Step 4: Convert Excel to CSVs
    convert_excel_to_csvs()
    
    # Step 5: Show structure
    list_extracted_structure()
    
    print("\n" + "=" * 60)
    print("Migration complete!")
    print("- Sunglasses: unchanged")
    print("- Old watch data: moved to legacy/")
    print("- New watch data: extracted and converted")
    print("=" * 60)


if __name__ == "__main__":
    main()
